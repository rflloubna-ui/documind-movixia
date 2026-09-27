"""
Module Export : génère un fichier Word (.docx) ou PDF à partir d'un texte
(résumé ou traduction), pour téléchargement direct depuis l'application.
"""
import io
from docx import Document


def generer_docx(titre, contenu):
    """Retourne un buffer BytesIO contenant un document Word prêt à télécharger."""
    doc = Document()
    doc.add_heading(titre, level=1)
    for paragraphe in contenu.split("\n"):
        if paragraphe.strip():
            doc.add_paragraph(paragraphe)
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def generer_pdf(titre, contenu):
    """Retourne les octets d'un PDF prêt à télécharger (nécessite fpdf2)."""
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.multi_cell(0, 10, titre)
    pdf.ln(4)
    pdf.set_font("Helvetica", size=12)
    for paragraphe in contenu.split("\n"):
        if paragraphe.strip():
            pdf.multi_cell(0, 8, paragraphe)
            pdf.ln(2)
    sortie = pdf.output()
    return bytes(sortie)
